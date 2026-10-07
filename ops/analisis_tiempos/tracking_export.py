# -*- coding: utf-8 -*-
"""
Exportacion a Excel/CSV del uso del tiempo de los ejecutivos (seguimiento operativo).

Por cada empleado y dia habil (horario laboral configurado, por defecto 06:00-20:00,
lunes a viernes) se calcula a partir de las posiciones GPS de Traccar:
  - hora de inicio y fin del recorrido, duracion de la jornada
  - kilometros recorridos (sin saltos imposibles de GPS)
  - paradas: donde, a que hora, cuanto duraron y si fueron en un aliado
    (lugar de trabajo registrado) o en otro lugar; tiempo y distancia entre paradas
  - tiempo en aliados, en otros lugares y en movimiento
  - visitas a aliados y tiempo promedio entre visitas
Dias incoherentes (viajes fuera de la zona de operacion, >anomaly_km km, GPS sin datos
suficientes) se EXCLUYEN de los indicadores y se listan aparte con el motivo.
"""
import csv
import io
import logging
import zipfile
from collections import defaultdict
from datetime import datetime, timedelta

import pytz

from app.models import Ally, ProximityVisit, User
from app.traccar import get_device_positions
from app.utils import haversine_distance, filter_positions_by_working_hours

logger = logging.getLogger(__name__)
COLOMBIA_TZ = pytz.timezone('America/Bogota')

DEFAULTS = {
    'min_stop_min': 10,        # parada = quieto >= 10 min
    'stop_radius_m': 80,       # radio para agrupar puntos de una misma parada
    'ally_match_m': 200,       # distancia maxima parada->aliado para contarla como visita
    'max_speed_kmh': 180,      # salto GPS imposible por encima de esto
    'max_accuracy_m': 250,     # puntos con peor precision se ignoran
    'zone_lat': 4.65, 'zone_lng': -74.10, 'zone_radius_km': 150,   # zona habitual (se recalcula por empleado)
    'anomaly_km': 600,         # mas de esto en un dia = incoherente
    'min_points': 5,           # menos puntos que esto = dia sin datos utiles
}


def _parse_time(p):
    ft = p.get('fixTime')
    if not ft:
        return None
    try:
        return datetime.fromisoformat(str(ft).replace('Z', '+00:00')).astimezone(COLOMBIA_TZ)
    except (ValueError, AttributeError):
        return None


def _fmt_hm(dt):
    return dt.strftime('%H:%M') if dt else ''


def _mins(a, b):
    return round((b - a).total_seconds() / 60.0, 1)


def _clean_day_points(points, prm):
    """Ordena, quita puntos imprecisos y saltos imposibles. Devuelve (puntos, saltos_descartados)."""
    pts = []
    for p in points:
        lat, lng = p.get('latitude'), p.get('longitude')
        t = _parse_time(p)
        if lat is None or lng is None or t is None:
            continue
        acc = p.get('accuracy')
        if acc and acc > prm['max_accuracy_m']:
            continue
        pts.append({'t': t, 'lat': float(lat), 'lng': float(lng),
                    'speed': (p.get('speed') or 0) * 1.852, 'address': p.get('address') or ''})
    pts.sort(key=lambda x: x['t'])
    out, jumps = [], 0
    for p in pts:
        if out:
            q = out[-1]
            dt_h = (p['t'] - q['t']).total_seconds() / 3600.0
            d_km = haversine_distance(q['lat'], q['lng'], p['lat'], p['lng']) / 1000.0
            if dt_h > 0 and d_km / dt_h > prm['max_speed_kmh'] and d_km > 0.5:
                jumps += 1
                continue
        out.append(p)
    return out, jumps


def _detect_stops(pts, prm):
    """Agrupa puntos consecutivos dentro de stop_radius_m; una parada dura >= min_stop_min."""
    stops = []
    i, n = 0, len(pts)
    while i < n:
        c_lat, c_lng, k = pts[i]['lat'], pts[i]['lng'], 1
        j = i + 1
        while j < n and haversine_distance(c_lat, c_lng, pts[j]['lat'], pts[j]['lng']) <= prm['stop_radius_m']:
            k += 1
            c_lat += (pts[j]['lat'] - c_lat) / k
            c_lng += (pts[j]['lng'] - c_lng) / k
            j += 1
        dur = _mins(pts[i]['t'], pts[j - 1]['t'])
        if dur >= prm['min_stop_min']:
            addr = next((p['address'] for p in pts[i:j] if p['address']), '')
            stops.append({'start': pts[i]['t'], 'end': pts[j - 1]['t'], 'min': dur,
                          'lat': round(c_lat, 6), 'lng': round(c_lng, 6), 'address': addr})
            i = j
        else:
            i += 1
    return stops


def _classify_stop(stop, allies, prm):
    best, best_d = None, None
    for a in allies:
        d = haversine_distance(stop['lat'], stop['lng'], a.latitude, a.longitude)
        lim = max(a.radius or 0, prm['ally_match_m'])
        if d <= lim and (best_d is None or d < best_d):
            best, best_d = a, d
    if best:
        return 'Aliado', best.name, round(best_d)
    return 'Otro lugar', '', None


def build_tracking_dataset(users, start_d, end_d, params=None, fetch_budget_s=80):
    """Devuelve dict con filas para cada hoja. users: lista de User con traccar_device_id."""
    from app.analytics.commercial import _parallel_fetch
    prm = dict(DEFAULTS)
    if params:
        prm.update({k: v for k, v in params.items() if v is not None})

    start_utc = COLOMBIA_TZ.localize(datetime.combine(start_d, datetime.min.time())).astimezone(pytz.utc)
    end_utc = COLOMBIA_TZ.localize(datetime.combine(end_d + timedelta(days=1), datetime.min.time())).astimezone(pytz.utc)
    allies = [a for a in Ally.query.all() if a.latitude is not None and a.longitude is not None]
    users = [u for u in users if u.traccar_device_id]
    pos_map = _parallel_fetch([u.traccar_device_id for u in users],
                              lambda did: get_device_positions(did, start_utc, end_utc),
                              budget_s=fetch_budget_s, max_workers=4)

    resumen, paradas, excluidos, por_empleado = [], [], [], []
    for u in users:
        name = u.full_name or u.username
        raw = pos_map.get(u.traccar_device_id)
        if raw is None:
            excluidos.append({'empleado': name, 'fecha': '', 'motivo': 'Traccar no respondio a tiempo (sin datos del periodo)', 'km': ''})
            continue
        work = filter_positions_by_working_hours(raw)
        by_day = defaultdict(list)
        for p in work:
            t = _parse_time(p)
            if t:
                by_day[t.date()].append(p)
        agg = {'dias': 0, 'km': 0.0, 'jornada_min': 0.0, 'paradas': 0, 'visitas': 0, 'min_aliados': 0.0,
               'min_otros': 0.0, 'min_mov': 0.0, 'inicios': [], 'fines': [], 'gaps': []}
        # Zona habitual DEL EMPLEADO: mediana de sus posiciones del periodo. Un dia lejos de
        # su zona (viaje a otra region) se excluye; asi vale para equipos de Bogota, Boyaca, etc.
        lats = sorted(float(p['latitude']) for p in work if p.get('latitude') is not None)
        lngs = sorted(float(p['longitude']) for p in work if p.get('longitude') is not None)
        if lats:
            prm = dict(prm, zone_lat=lats[len(lats) // 2], zone_lng=lngs[len(lngs) // 2])
        for day in sorted(by_day):
            if day.weekday() >= 5:
                continue
            pts, jumps = _clean_day_points(by_day[day], prm)
            if len(pts) < prm['min_points']:
                excluidos.append({'empleado': name, 'fecha': day.isoformat(), 'motivo': f'Sin datos GPS suficientes ({len(pts)} puntos)', 'km': ''})
                continue
            # Zona de operacion: si la mayoria del dia esta lejos (viaje a otra region) se excluye
            far = sum(1 for p in pts if haversine_distance(p['lat'], p['lng'], prm['zone_lat'], prm['zone_lng']) > prm['zone_radius_km'] * 1000)
            if far / len(pts) > 0.5:
                excluidos.append({'empleado': name, 'fecha': day.isoformat(), 'motivo': f'Viaje fuera de su zona habitual (>{prm["zone_radius_km"]} km, otra region)', 'km': ''})
                continue
            if far:
                pts = [p for p in pts if haversine_distance(p['lat'], p['lng'], prm['zone_lat'], prm['zone_lng']) <= prm['zone_radius_km'] * 1000]
            km = sum(haversine_distance(pts[i - 1]['lat'], pts[i - 1]['lng'], pts[i]['lat'], pts[i]['lng']) for i in range(1, len(pts))) / 1000.0
            if km > prm['anomaly_km']:
                excluidos.append({'empleado': name, 'fecha': day.isoformat(), 'motivo': f'Kilometraje incoherente ({km:.0f} km en un dia)', 'km': round(km, 1)})
                continue
            stops = _detect_stops(pts, prm)
            inicio, fin = pts[0]['t'], pts[-1]['t']
            jornada = _mins(inicio, fin)
            min_al = min_ot = 0.0
            n_vis = 0
            prev_end, prev_latlng, prev_ally_end = None, None, None
            gaps = []
            for k, s in enumerate(stops, start=1):
                tipo, aliado, dist = _classify_stop(s, allies, prm)
                if tipo == 'Aliado':
                    min_al += s['min']; n_vis += 1
                    if prev_ally_end:
                        gaps.append(_mins(prev_ally_end, s['start']))
                    prev_ally_end = s['end']
                else:
                    min_ot += s['min']
                travel_min = _mins(prev_end, s['start']) if prev_end else _mins(inicio, s['start'])
                travel_km = (haversine_distance(prev_latlng[0], prev_latlng[1], s['lat'], s['lng']) / 1000.0) if prev_latlng else None
                paradas.append({
                    'empleado': name, 'mes': day.strftime('%Y-%m'), 'fecha': day.isoformat(), 'dia': ['Lun', 'Mar', 'Mie', 'Jue', 'Vie', 'Sab', 'Dom'][day.weekday()],
                    'parada_n': k, 'llegada': _fmt_hm(s['start']), 'salida': _fmt_hm(s['end']), 'duracion_min': s['min'],
                    'tipo': tipo, 'aliado': aliado, 'dist_al_aliado_m': dist if dist is not None else '',
                    'latitud': s['lat'], 'longitud': s['lng'], 'direccion': s['address'],
                    'desplazamiento_min_desde_anterior': travel_min, 'km_desde_anterior': round(travel_km, 2) if travel_km is not None else '',
                    'mapa': f"https://www.google.com/maps?q={s['lat']},{s['lng']}",
                })
                prev_end, prev_latlng = s['end'], (s['lat'], s['lng'])
            min_mov = max(0.0, jornada - min_al - min_ot)
            resumen.append({
                'empleado': name, 'mes': day.strftime('%Y-%m'), 'fecha': day.isoformat(), 'dia': ['Lun', 'Mar', 'Mie', 'Jue', 'Vie', 'Sab', 'Dom'][day.weekday()],
                'inicio_recorrido': _fmt_hm(inicio), 'fin_recorrido': _fmt_hm(fin), 'jornada_horas': round(jornada / 60.0, 2),
                'km_recorridos': round(km, 1), 'paradas': len(stops), 'visitas_aliados': n_vis,
                'min_en_aliados': round(min_al), 'min_en_otros_lugares': round(min_ot), 'min_en_movimiento': round(min_mov),
                'pct_tiempo_aliados': round(100 * min_al / jornada, 1) if jornada else 0,
                'pct_tiempo_otros': round(100 * min_ot / jornada, 1) if jornada else 0,
                'pct_tiempo_movimiento': round(100 * min_mov / jornada, 1) if jornada else 0,
                'tiempo_prom_entre_visitas_min': round(sum(gaps) / len(gaps)) if gaps else '',
                'vel_max_kmh': round(max(p['speed'] for p in pts)), 'puntos_gps': len(pts), 'saltos_gps_descartados': jumps,
            })
            agg['dias'] += 1; agg['km'] += km; agg['jornada_min'] += jornada; agg['paradas'] += len(stops)
            agg['visitas'] += n_vis; agg['min_aliados'] += min_al; agg['min_otros'] += min_ot; agg['min_mov'] += min_mov
            agg['inicios'].append(inicio.hour * 60 + inicio.minute); agg['fines'].append(fin.hour * 60 + fin.minute); agg['gaps'] += gaps
        d = agg['dias'] or 1
        avg_hm = lambda lst: (f"{int(sum(lst) / len(lst)) // 60:02d}:{int(sum(lst) / len(lst)) % 60:02d}" if lst else '')
        tj = agg['jornada_min'] or 1
        por_empleado.append({
            'empleado': name, 'dias_con_datos': agg['dias'], 'km_total': round(agg['km'], 1), 'km_promedio_dia': round(agg['km'] / d, 1),
            'inicio_promedio': avg_hm(agg['inicios']), 'fin_promedio': avg_hm(agg['fines']), 'jornada_promedio_horas': round(agg['jornada_min'] / d / 60.0, 2),
            'paradas_promedio_dia': round(agg['paradas'] / d, 1), 'visitas_aliados_total': agg['visitas'], 'visitas_aliados_prom_dia': round(agg['visitas'] / d, 2),
            'pct_tiempo_aliados': round(100 * agg['min_aliados'] / tj, 1), 'pct_tiempo_otros_lugares': round(100 * agg['min_otros'] / tj, 1),
            'pct_tiempo_movimiento': round(100 * agg['min_mov'] / tj, 1),
            'tiempo_prom_entre_visitas_min': round(sum(agg['gaps']) / len(agg['gaps'])) if agg['gaps'] else '',
            'dias_excluidos': sum(1 for e in excluidos if e['empleado'] == name and e['fecha']),
        })

    # Visitas por proximidad (tabla reprocesada en segundo plano): 1 por aliado y dia
    ally_names = {a.id: a.name for a in Ally.query.all()}
    uid_name = {u.id: (u.full_name or u.username) for u in users}
    visitas = []
    for v in ProximityVisit.query.filter(ProximityVisit.visit_date >= start_d, ProximityVisit.visit_date <= end_d,
                                         ProximityVisit.user_id.in_([u.id for u in users]) if users else False).all():
        if v.visit_date.weekday() >= 5:
            continue
        ft = pytz.utc.localize(v.first_time).astimezone(COLOMBIA_TZ) if v.first_time else None
        lt = pytz.utc.localize(v.last_time).astimezone(COLOMBIA_TZ) if v.last_time else None
        visitas.append({'empleado': uid_name.get(v.user_id, v.user_id), 'fecha': v.visit_date.isoformat(), 'aliado': ally_names.get(v.ally_id, v.ally_id),
                        'primera_hora_en_radio': _fmt_hm(ft), 'ultima_hora_en_radio': _fmt_hm(lt),
                        'permanencia_min': v.duration_minutes if hasattr(v, 'duration_minutes') else '', 'radio_m': v.radius_m})
    visitas.sort(key=lambda r: (r['empleado'], r['fecha'], r['primera_hora_en_radio']))

    parametros = [
        {'parametro': 'Periodo', 'valor': f'{start_d.isoformat()} a {end_d.isoformat()}'},
        {'parametro': 'Horario laboral', 'valor': 'Segun configuracion (por defecto 06:00-20:00, lunes a viernes). Sabados y domingos excluidos.'},
        {'parametro': 'Parada', 'valor': f">= {prm['min_stop_min']} min quieto dentro de {prm['stop_radius_m']} m"},
        {'parametro': 'Visita a aliado', 'valor': f"parada a <= {prm['ally_match_m']} m (o el radio del aliado) de su ubicacion registrada"},
        {'parametro': 'Limpieza GPS', 'valor': f"se descartan puntos con precision > {prm['max_accuracy_m']} m y saltos > {prm['max_speed_kmh']} km/h"},
        {'parametro': 'Dias excluidos', 'valor': f"viaje fuera de la zona habitual del empleado (> {prm['zone_radius_km']} km de su centro de operacion la mayor parte del dia), > {prm['anomaly_km']} km/dia, o < {prm['min_points']} puntos GPS"},
        {'parametro': 'Empleados', 'valor': ', '.join(u.full_name or u.username for u in users)},
        {'parametro': 'Generado', 'valor': datetime.now(COLOMBIA_TZ).strftime('%Y-%m-%d %H:%M')},
    ]
    lugares = _frequent_places(paradas, resumen)
    _add_base_times(paradas, resumen)
    if prm.get('reverse_geocode'):
        _reverse_geocode_places(lugares, paradas)
    return {'Resumen por empleado': por_empleado, 'Resumen diario': resumen, 'Paradas': paradas,
            'Lugares frecuentes': lugares, 'Visitas aliados': visitas, 'Dias excluidos': excluidos, 'Parametros': parametros}


BASES = {}   # empleado -> (lat, lng) de su base/domicilio detectado


def _frequent_places(paradas, resumen, radius_m=150, min_times=3):
    """Agrupa las paradas de cada empleado por lugar (<= radius_m) para ver en que sitios
    repite: cuantas veces, cuantos dias, hora promedio de llegada y duracion. Marca el
    posible domicilio (lugar donde arranca o termina el dia con frecuencia)."""
    out = []
    by_emp = defaultdict(list)
    for p in paradas:
        by_emp[p['empleado']].append(p)
    first_last = defaultdict(set)   # (empleado, fecha) -> parada_n de la primera y la ultima
    for p in paradas:
        key = (p['empleado'], p['fecha'])
        first_last[key].add(p['parada_n'])
    for emp, rows in by_emp.items():
        clusters = []   # [lat, lng, [rows]]
        for p in rows:
            for c in clusters:
                if haversine_distance(c[0], c[1], p['latitud'], p['longitud']) <= radius_m:
                    n = len(c[2]); c[0] += (p['latitud'] - c[0]) / (n + 1); c[1] += (p['longitud'] - c[1]) / (n + 1); c[2].append(p); break
            else:
                clusters.append([p['latitud'], p['longitud'], [p]])
        clusters.sort(key=lambda c: -len(c[2]))
        days_emp = {r['fecha'] for r in rows}
        # Base/domicilio: el lugar (sin aliado) donde mas veces empieza o termina el dia,
        # si eso ocurre en al menos el 40% de los dias con datos. Solo uno por empleado.
        def score(c):
            rs = c[2]
            if any(r['aliado'] for r in rs) or len(rs) < min_times:
                return -1
            return sum(1 for r in rs if r['parada_n'] == 1 or r['parada_n'] == max(first_last[(emp, r['fecha'])]))
        best = max(clusters, key=score) if clusters else None
        base = best if best is not None and score(best) >= max(3, 0.4 * len(days_emp)) else None
        if base is not None:
            BASES[emp] = (base[0], base[1])
        for idx, c in enumerate(clusters, start=1):
            rs = c[2]
            if len(rs) < min_times:
                continue
            label = f'Lugar frecuente #{idx}'
            dias = {r['fecha'] for r in rs}
            firsts = sum(1 for r in rs if r['parada_n'] == 1)
            lasts = sum(1 for r in rs if r['parada_n'] == max(first_last[(emp, r['fecha'])]))
            ally = next((r['aliado'] for r in rs if r['aliado']), '')
            posible = 'Posible domicilio / base' if c is base else ('' if ally else 'Lugar recurrente')
            mins = [int(r['llegada'][:2]) * 60 + int(r['llegada'][3:]) for r in rs if r['llegada']]
            avg = f"{(sum(mins) // len(mins)) // 60:02d}:{(sum(mins) // len(mins)) % 60:02d}" if mins else ''
            for r in rs:
                r['lugar_frecuente'] = label + (' · ' + posible if posible else '')
            out.append({'empleado': emp, 'lugar': label, 'tipo': ('Aliado: ' + ally) if ally else posible,
                        'veces': len(rs), 'dias_distintos': len(dias), 'llegada_promedio': avg,
                        'duracion_promedio_min': round(sum(r['duracion_min'] for r in rs) / len(rs)),
                        'minutos_totales': round(sum(r['duracion_min'] for r in rs)),
                        'veces_primera_parada_del_dia': firsts, 'veces_ultima_parada_del_dia': lasts,
                        'latitud': round(c[0], 6), 'longitud': round(c[1], 6), 'direccion': next((r['direccion'] for r in rs if r['direccion']), ''),
                        'mapa': f'https://www.google.com/maps?q={round(c[0], 6)},{round(c[1], 6)}'})
    for p in paradas:
        p.setdefault('lugar_frecuente', '')
    return out


def _add_base_times(paradas, resumen):
    """Hora en que el empleado SALE de su base/domicilio (fin de la primera parada si es la base)
    y hora en que LLEGA de vuelta (inicio de la ultima parada si es la base)."""
    by_day = defaultdict(list)
    for p in paradas:
        by_day[(p['empleado'], p['fecha'])].append(p)
    for r in resumen:
        base = BASES.get(r['empleado'])
        r['salida_de_base'] = r['llegada_a_base'] = ''
        r['horas_fuera_de_base'] = ''
        if not base:
            continue
        ps = sorted(by_day.get((r['empleado'], r['fecha']), []), key=lambda x: x['parada_n'])
        if not ps:
            continue
        first, last = ps[0], ps[-1]
        near = lambda p: haversine_distance(p['latitud'], p['longitud'], base[0], base[1]) <= 150
        if near(first):
            r['salida_de_base'] = first['salida']
        if near(last) and (len(ps) > 1 or not near(first)):
            r['llegada_a_base'] = last['llegada']
        if r['salida_de_base'] and r['llegada_a_base']:
            a = int(r['salida_de_base'][:2]) * 60 + int(r['salida_de_base'][3:])
            b = int(r['llegada_a_base'][:2]) * 60 + int(r['llegada_a_base'][3:])
            r['horas_fuera_de_base'] = round(max(0, b - a) / 60.0, 2)


def _reverse_geocode_places(lugares, paradas):
    """Direccion aproximada de los lugares frecuentes (Nominatim, 1 consulta/s)."""
    import time, requests
    cache = {}
    for l in lugares:
        key = (round(l['latitud'], 4), round(l['longitud'], 4))
        if key not in cache:
            try:
                r = requests.get('https://nominatim.openstreetmap.org/reverse',
                                 params={'lat': l['latitud'], 'lon': l['longitud'], 'format': 'json', 'zoom': 17},
                                 headers={'User-Agent': 'VantiGo-Analisis/1.0 (contacto: bgaleanotec@gmail.com)'}, timeout=8)
                cache[key] = r.json().get('display_name', '') if r.ok else ''
            except Exception:
                cache[key] = ''
            time.sleep(1.1)
        l['direccion'] = cache[key]
    for p in paradas:
        if not p.get('direccion') and p.get('lugar_frecuente'):
            for l in lugares:
                if l['empleado'] == p['empleado'] and haversine_distance(l['latitud'], l['longitud'], p['latitud'], p['longitud']) <= 150:
                    p['direccion'] = l['direccion']
                    break


HEADERS_ES = {
    'empleado': 'Empleado', 'mes': 'Mes', 'fecha': 'Fecha', 'dia': 'Dia', 'inicio_recorrido': 'Inicio recorrido', 'fin_recorrido': 'Fin recorrido',
    'jornada_horas': 'Jornada (h)', 'km_recorridos': 'Km recorridos', 'paradas': 'Paradas', 'visitas_aliados': 'Visitas a aliados',
    'min_en_aliados': 'Min en aliados', 'min_en_otros_lugares': 'Min en otros lugares', 'min_en_movimiento': 'Min en movimiento',
    'pct_tiempo_aliados': '% tiempo en aliados', 'pct_tiempo_otros': '% tiempo otros lugares', 'pct_tiempo_movimiento': '% tiempo en movimiento',
    'tiempo_prom_entre_visitas_min': 'Tiempo prom. entre visitas (min)', 'vel_max_kmh': 'Vel. max (km/h)', 'puntos_gps': 'Puntos GPS',
    'saltos_gps_descartados': 'Saltos GPS descartados', 'parada_n': 'Parada #', 'llegada': 'Llegada', 'salida': 'Salida', 'duracion_min': 'Duracion (min)',
    'tipo': 'Tipo de lugar', 'aliado': 'Aliado', 'dist_al_aliado_m': 'Dist. al aliado (m)', 'latitud': 'Latitud', 'longitud': 'Longitud', 'direccion': 'Direccion',
    'desplazamiento_min_desde_anterior': 'Desplazamiento desde anterior (min)', 'km_desde_anterior': 'Km desde anterior', 'mapa': 'Ver en mapa',
    'dias_con_datos': 'Dias con datos', 'km_total': 'Km total', 'km_promedio_dia': 'Km promedio/dia', 'inicio_promedio': 'Inicio promedio', 'fin_promedio': 'Fin promedio',
    'jornada_promedio_horas': 'Jornada promedio (h)', 'paradas_promedio_dia': 'Paradas promedio/dia', 'visitas_aliados_total': 'Visitas aliados (total)',
    'visitas_aliados_prom_dia': 'Visitas aliados/dia', 'pct_tiempo_otros_lugares': '% tiempo otros lugares', 'dias_excluidos': 'Dias excluidos',
    'primera_hora_en_radio': 'Primera hora en radio', 'ultima_hora_en_radio': 'Ultima hora en radio', 'permanencia_min': 'Permanencia (min)', 'radio_m': 'Radio (m)',
    'motivo': 'Motivo', 'km': 'Km', 'parametro': 'Parametro', 'valor': 'Valor',
    'lugar_frecuente': 'Lugar frecuente', 'lugar': 'Lugar', 'veces': 'Veces', 'dias_distintos': 'Dias distintos', 'llegada_promedio': 'Llegada promedio',
    'duracion_promedio_min': 'Duracion promedio (min)', 'minutos_totales': 'Minutos totales', 'veces_primera_parada_del_dia': 'Veces 1ra parada del dia',
    'veces_ultima_parada_del_dia': 'Veces ultima parada del dia', 'salida_de_base': 'Salida de base/domicilio',
    'llegada_a_base': 'Llegada a base/domicilio', 'horas_fuera_de_base': 'Horas fuera de base',
}


def _rows_to_table(rows):
    if not rows:
        return [], []
    keys = list(rows[0].keys())
    return [HEADERS_ES.get(k, k) for k in keys], [[r.get(k, '') for k in keys] for r in rows]


def to_xlsx(dataset):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); wb.remove(wb.active)
    head_fill = PatternFill('solid', fgColor='4F46E5'); head_font = Font(bold=True, color='FFFFFF')
    for sheet, rows in dataset.items():
        ws = wb.create_sheet(sheet[:31])
        headers, data = _rows_to_table(rows)
        if not headers:
            ws.append(['Sin datos en el periodo']); continue
        ws.append(headers)
        for c in ws[1]:
            c.fill = head_fill; c.font = head_font; c.alignment = Alignment(vertical='center', wrap_text=True)
        for r in data:
            ws.append(r)
        ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
        for i, h in enumerate(headers, start=1):
            width = max(len(str(h)), *(len(str(r[i - 1])) for r in data[:200])) if data else len(str(h))
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 48)
        if sheet == 'Paradas':
            col = headers.index('Ver en mapa') + 1
            for row in range(2, len(data) + 2):
                cell = ws.cell(row=row, column=col); cell.hyperlink = cell.value; cell.font = Font(color='2563EB', underline='single')
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf.getvalue()


def to_csv_zip(dataset):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for sheet, rows in dataset.items():
            headers, data = _rows_to_table(rows)
            s = io.StringIO(); w = csv.writer(s, delimiter=';', lineterminator='\n')
            w.writerow(headers or ['Sin datos en el periodo'])
            for r in data:
                w.writerow(r)
            z.writestr(sheet.replace(' ', '_').lower() + '.csv', '﻿' + s.getvalue())   # BOM: Excel abre con acentos
    buf.seek(0)
    return buf.getvalue()
