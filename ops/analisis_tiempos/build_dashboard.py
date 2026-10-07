# -*- coding: utf-8 -*-
"""Construye el Super Dashboard (HTML interactivo) y la hoja 'Dashboard' con graficos
dentro del Excel, a partir del libro generado por run_analisis.py.
Uso: python ops/analisis_tiempos/build_dashboard.py <ruta.xlsx>"""
import io, json, os, sys
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill, Alignment

XLSX = sys.argv[1] if len(sys.argv) > 1 else os.path.join(os.path.expanduser('~'), 'OneDrive', 'Desktop', 'Analisis_Tiempos_Ejecutivos_2026-01-01_a_2026-10-06.xlsx')

# Segmentos: Mercado Comercial = registro antiguo; Residencial = bloque que llego el 20-27 de agosto de 2026
RESIDENCIAL = {'amgomez', 'jjbecerra', 'kvlamprea', 'vqherrera', 'gmalvarado', 'elmateus', 'lemurcia', 'dramaya', 'malejandra', 'kmlara', 'ajsuarez'}
SEG = lambda e: 'Residencial' if e in RESIDENCIAL else 'Mercado Comercial'

wb = load_workbook(XLSX, read_only=True)
def sheet(name):
    it = wb[name].iter_rows(values_only=True); h = next(it); return [dict(zip(h, r)) for r in it if any(v not in (None, '') for v in r)]
diario, paradas, lugares, excl, emp = sheet('Resumen diario'), sheet('Paradas'), sheet('Lugares frecuentes'), sheet('Dias excluidos'), sheet('Resumen por empleado')
wb.close()

def num(v):
    try: return float(v) if v not in (None, '') else None
    except (TypeError, ValueError): return None

data = {
    'segmentos': {e['Empleado']: SEG(e['Empleado']) for e in emp},
    'diario': [{'e': r['Empleado'], 'f': r['Fecha'], 'd': r['Dia'], 'ini': r['Inicio recorrido'], 'fin': r['Fin recorrido'], 'h': num(r['Jornada (h)']),
                'km': num(r['Km recorridos']), 'p': num(r['Paradas']), 'mL': num(r['Min en aliados']) or 0, 'mO': num(r['Min en otros lugares']) or 0,
                'mM': num(r['Min en movimiento']) or 0, 'vmax': num(r['Vel. max (km/h)']), 'sb': r.get('Salida de base/domicilio') or '', 'lb': r.get('Llegada a base/domicilio') or '',
                'hf': num(r.get('Horas fuera de base'))} for r in diario],
    'paradas': [{'e': r['Empleado'], 'f': r['Fecha'], 'n': r['Parada #'], 'll': r['Llegada'], 'sa': r['Salida'], 'min': num(r['Duracion (min)']), 't': r['Tipo de lugar'],
                 'al': r.get('Aliado') or '', 'lat': num(r['Latitud']), 'lng': num(r['Longitud']), 'dir': r.get('Direccion') or '', 'lf': r.get('Lugar frecuente') or '',
                 'tr': num(r['Desplazamiento desde anterior (min)']), 'kma': num(r['Km desde anterior'])} for r in paradas],
    'lugares': [{'e': r['Empleado'], 'l': r['Lugar'], 't': r.get('Tipo') or r.get('Tipo de lugar') or '', 'v': num(r['Veces']), 'dd': num(r['Dias distintos']), 'll': r['Llegada promedio'],
                 'dm': num(r['Duracion promedio (min)']), 'mt': num(r['Minutos totales']), 'lat': num(r['Latitud']), 'lng': num(r['Longitud']), 'dir': r.get('Direccion') or ''} for r in lugares],
    'excluidos': [{'e': r['Empleado'], 'f': r['Fecha'], 'm': r['Motivo']} for r in excl],
    'periodo': [min(r['Fecha'] for r in diario), max(r['Fecha'] for r in diario)],
}

HTML = r'''<!DOCTYPE html>
<html lang="es"><head><meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>VantiGo · Super Dashboard de Tiempos</title>
<link rel="stylesheet" href="https://unpkg.com/leaflet@1.9.4/dist/leaflet.css">
<script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.1/dist/chart.umd.min.js"></script>
<script src="https://unpkg.com/leaflet@1.9.4/dist/leaflet.js"></script>
<style>
:root{--bg:#0b1020;--card:#121a30;--card2:#1a2340;--line:#243052;--ind:#6366f1;--ind2:#818cf8;--em:#10b981;--am:#f59e0b;--red:#ef4444;--cy:#22d3ee;--tx:#f1f5f9;--mu:#94a3b8}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--tx);font-family:Inter,Segoe UI,system-ui,sans-serif;font-size:14px}
.wrap{max-width:1500px;margin:0 auto;padding:16px}
h1{font-size:24px;margin:0}h2{font-size:15px;margin:0 0 10px;color:var(--mu);text-transform:uppercase;letter-spacing:.5px}
.top{display:flex;flex-wrap:wrap;gap:12px;align-items:center;justify-content:space-between;margin-bottom:14px}
.badge{display:inline-block;padding:3px 10px;border-radius:999px;font-size:11px;font-weight:700}
.b-com{background:rgba(99,102,241,.2);color:#a5b4fc;border:1px solid rgba(99,102,241,.4)}.b-res{background:rgba(16,185,129,.18);color:#6ee7b7;border:1px solid rgba(16,185,129,.4)}
.filters{display:flex;flex-wrap:wrap;gap:8px;align-items:center;background:var(--card);border:1px solid var(--line);border-radius:14px;padding:10px 12px;margin-bottom:14px}
.filters label{font-size:11px;color:var(--mu)}select,input[type=date]{background:#0f172a;color:var(--tx);border:1px solid var(--line);border-radius:10px;padding:7px 9px;font-size:13px}
.chip{cursor:pointer;padding:6px 12px;border-radius:999px;border:1px solid var(--line);background:#0f172a;color:var(--mu);font-size:12px;font-weight:700}.chip.on{background:var(--ind);color:#fff;border-color:var(--ind)}
.btn{cursor:pointer;padding:7px 12px;border-radius:10px;border:1px solid var(--line);background:#0f172a;color:var(--tx);font-size:12px;font-weight:700}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px;margin-bottom:14px}
.kpi{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:12px 14px}.kpi .v{font-size:26px;font-weight:800}.kpi .l{font-size:11px;color:var(--mu);margin-top:2px}.kpi .s{font-size:11px;margin-top:4px}
.grid{display:grid;grid-template-columns:repeat(12,1fr);gap:12px}.card{background:var(--card);border:1px solid var(--line);border-radius:14px;padding:14px;grid-column:span 12}
@media(min-width:900px){.c6{grid-column:span 6}.c4{grid-column:span 4}.c8{grid-column:span 8}}
.chart{position:relative;height:300px}.chart.tall{height:420px}
table{width:100%;border-collapse:collapse;font-size:12.5px}th{text-align:left;color:var(--mu);font-size:11px;text-transform:uppercase;padding:8px 6px;border-bottom:1px solid var(--line);cursor:pointer;white-space:nowrap}
td{padding:7px 6px;border-bottom:1px solid rgba(36,48,82,.5);white-space:nowrap}tr:hover td{background:rgba(99,102,241,.06)}td.r,th.r{text-align:right}
.bar{height:6px;border-radius:3px;background:var(--ind);display:inline-block;vertical-align:middle;margin-left:6px}
.seg{display:grid;grid-template-columns:1fr 1fr;gap:12px}@media(max-width:800px){.seg{grid-template-columns:1fr}}
.segcard{border-radius:14px;padding:14px;border:1px solid var(--line)}.segcard.com{background:linear-gradient(135deg,rgba(99,102,241,.15),transparent)}.segcard.res{background:linear-gradient(135deg,rgba(16,185,129,.14),transparent)}
.segcard .row{display:flex;justify-content:space-between;padding:5px 0;border-bottom:1px dashed rgba(148,163,184,.15);font-size:13px}.segcard .row b{font-size:14px}
#map{height:380px;border-radius:12px}.leaflet-container{background:#dce8d5}
.heat{display:grid;grid-template-columns:70px repeat(14,1fr);gap:3px;font-size:10px}.heat div{padding:4px 2px;text-align:center;border-radius:4px;color:#e2e8f0}
.tl{display:flex;flex-direction:column;gap:6px}.tl .it{display:grid;grid-template-columns:52px 1fr;gap:8px;align-items:start}.tl .t{color:var(--ind2);font-weight:700;font-size:12px}
.tl .box{background:var(--card2);border-radius:10px;padding:8px 10px;font-size:12px}.tl .mv{color:var(--mu);font-size:11px;padding:2px 10px}
.foot{color:var(--mu);font-size:11px;margin:18px 0 6px;text-align:center}.pill{font-size:10px;padding:2px 7px;border-radius:999px;background:#0f172a;border:1px solid var(--line);color:var(--mu)}
.note{font-size:11px;color:var(--mu);margin-top:6px}
</style></head><body><div class="wrap">
<div class="top"><div><h1>Super Dashboard · Uso del tiempo de los ejecutivos</h1><div class="note">VantiGo · Horario laboral 06:00–20:00, lunes a viernes · Datos GPS depurados (sin saltos ni viajes fuera de zona) · Periodo disponible: <b id="per"></b></div></div>
<div><span class="badge b-com">Mercado Comercial</span> <span class="badge b-res">Residencial</span></div></div>

<div class="filters">
  <span class="chip on" data-seg="Todos">Todos</span><span class="chip" data-seg="Mercado Comercial">Mercado Comercial</span><span class="chip" data-seg="Residencial">Residencial</span>
  <label>Ejecutivo</label><select id="fEmp"><option value="">Todos</option></select>
  <label>Mes</label><select id="fMes"><option value="">Todos</option></select>
  <label>Desde</label><input type="date" id="fIni"><label>Hasta</label><input type="date" id="fFin">
  <button class="btn" id="fReset">Limpiar</button><span class="note" id="fInfo"></span>
</div>

<div class="kpis" id="kpis"></div>

<div class="grid">
  <div class="card"><h2>Comparativo de segmentos</h2><div class="seg" id="segcards"></div><div class="note">Mercado Comercial: ejecutivos con registro antiguo (desde enero–mayo). Residencial: bloque que ingreso el 20–27 de agosto. Los promedios son por dia con datos, dentro del horario laboral.</div></div>
  <div class="card c6"><h2>Kilometros por mes (promedio por dia y por ejecutivo)</h2><div class="chart"><canvas id="cMes"></canvas></div></div>
  <div class="card c6"><h2>Dias con actividad por mes</h2><div class="chart"><canvas id="cDias"></canvas></div></div>
  <div class="card c6"><h2>Jornada: inicio y fin promedio por ejecutivo</h2><div class="chart tall"><canvas id="cJornada"></canvas></div></div>
  <div class="card c6"><h2>Distribucion del tiempo por ejecutivo</h2><div class="chart tall"><canvas id="cTiempo"></canvas></div><div class="note">Detenido = paradas de 10 min o mas (en aliado u otro lugar). En movimiento = resto de la jornada.</div></div>
  <div class="card"><h2>Ranking de ejecutivos <span class="pill">clic en el encabezado para ordenar · clic en la fila para ver el detalle</span></h2><div style="overflow-x:auto"><table id="rank"><thead><tr>
    <th data-k="e">Ejecutivo</th><th data-k="seg">Segmento</th><th class="r" data-k="dias">Dias</th><th class="r" data-k="kmd">Km/dia</th><th class="r" data-k="km">Km total</th><th class="r" data-k="ini">Inicio prom.</th><th class="r" data-k="fin">Fin prom.</th><th class="r" data-k="h">Jornada (h)</th><th class="r" data-k="pd">Paradas/dia</th><th class="r" data-k="pl">% detenido</th><th class="r" data-k="pm">% movimiento</th><th class="r" data-k="hf">Horas fuera de base</th><th class="r" data-k="ex">Dias excluidos</th>
  </tr></thead><tbody></tbody></table></div></div>
  <div class="card c6"><h2>Mapa de calor: a que horas y que dias paran</h2><div class="heat" id="heat"></div><div class="note">Cantidad de paradas (de 10 min o mas) por dia de la semana y hora de llegada.</div></div>
  <div class="card c6"><h2>Paradas: duracion tipica</h2><div class="chart"><canvas id="cDur"></canvas></div></div>
  <div class="card c8"><h2>Mapa de lugares frecuentes <span class="pill">tamano = minutos acumulados · color = segmento · rojo = posible base/domicilio</span></h2><div id="map"></div></div>
  <div class="card c4"><h2>Lugares frecuentes</h2><div style="overflow:auto;max-height:380px"><table id="tLug"><thead><tr><th>Ejecutivo</th><th>Tipo</th><th class="r">Veces</th><th class="r">Min tot.</th><th>Llegada</th></tr></thead><tbody></tbody></table></div></div>
  <div class="card"><h2>Detalle del ejecutivo: <span id="detName" style="color:var(--ind2)">(selecciona uno en el ranking o en el filtro)</span></h2>
    <div class="grid" style="margin-top:6px"><div class="card c8" style="background:var(--card2)"><h2>Calendario de kilometros por dia <span class="pill">clic en un dia para ver su recorrido</span></h2><div class="heat" id="cal" style="grid-template-columns:repeat(auto-fill,minmax(44px,1fr))"></div></div>
    <div class="card c4" style="background:var(--card2)"><h2>Recorrido del dia <span id="detDay"></span></h2><div class="tl" id="tl"><div class="note">Elige un dia en el calendario.</div></div></div></div></div>
  <div class="card"><h2>Dias excluidos del analisis</h2><div style="overflow:auto;max-height:260px"><table id="tEx"><thead><tr><th>Ejecutivo</th><th>Fecha</th><th>Motivo</th></tr></thead><tbody></tbody></table></div></div>
</div>
<div class="foot">VantiGo · Estrategia comercial, planificacion y conduccion segura · Generado __FECHA__</div></div>
<script>
const D = __DATA__;
const SEGC = {'Mercado Comercial':'#6366f1','Residencial':'#10b981'};
const st = {seg:'Todos', emp:'', mes:'', ini:'', fin:'', sortK:'kmd', sortD:-1, det:'', day:''};
const emps = Object.keys(D.segmentos).sort(); const meses = [...new Set(D.diario.map(r=>r.f.slice(0,7)))].sort();
const $ = s => document.querySelector(s); const fmt = (n,d=1)=> n==null?'—':Number(n).toLocaleString('es-CO',{maximumFractionDigits:d});
const toMin = h => h ? parseInt(h.slice(0,2))*60+parseInt(h.slice(3,5)) : null; const toHM = m => m==null?'—':`${String(Math.floor(m/60)).padStart(2,'0')}:${String(Math.round(m%60)).padStart(2,'0')}`;
const avg = a => a.length? a.reduce((x,y)=>x+y,0)/a.length : null;
document.getElementById('per').textContent = D.periodo[0]+' a '+D.periodo[1];
emps.forEach(e=>{const o=document.createElement('option');o.value=e;o.textContent=e+' ('+(D.segmentos[e]==='Residencial'?'Res':'Com')+')';$('#fEmp').appendChild(o)});
meses.forEach(m=>{const o=document.createElement('option');o.value=m;o.textContent=m;$('#fMes').appendChild(o)});
document.querySelectorAll('.chip').forEach(c=>c.onclick=()=>{document.querySelectorAll('.chip').forEach(x=>x.classList.remove('on'));c.classList.add('on');st.seg=c.dataset.seg;render()});
$('#fEmp').onchange=e=>{st.emp=e.target.value;st.det=st.emp;render()};$('#fMes').onchange=e=>{st.mes=e.target.value;render()};
$('#fIni').onchange=e=>{st.ini=e.target.value;render()};$('#fFin').onchange=e=>{st.fin=e.target.value;render()};
$('#fReset').onclick=()=>{st.seg='Todos';st.emp='';st.mes='';st.ini='';st.fin='';$('#fEmp').value='';$('#fMes').value='';$('#fIni').value='';$('#fFin').value='';document.querySelectorAll('.chip').forEach(x=>x.classList.toggle('on',x.dataset.seg==='Todos'));render()};
function keep(r){ if(st.seg!=='Todos'&&D.segmentos[r.e]!==st.seg)return false; if(st.emp&&r.e!==st.emp)return false; if(st.mes&&!r.f.startsWith(st.mes))return false; if(st.ini&&r.f<st.ini)return false; if(st.fin&&r.f>st.fin)return false; return true; }
const charts={}; function mk(id,cfg){ if(charts[id])charts[id].destroy(); charts[id]=new Chart(document.getElementById(id),cfg); }
Chart.defaults.color='#94a3b8'; Chart.defaults.borderColor='rgba(36,48,82,.6)'; Chart.defaults.font.family='Inter,Segoe UI,system-ui,sans-serif';
function perEmp(rows){ const m={}; rows.forEach(r=>{const o=m[r.e]||(m[r.e]={e:r.e,seg:D.segmentos[r.e],dias:0,km:0,h:0,p:0,mL:0,mO:0,mM:0,ini:[],fin:[],hf:[]}); o.dias++;o.km+=r.km||0;o.h+=r.h||0;o.p+=r.p||0;o.mL+=r.mL;o.mO+=r.mO;o.mM+=r.mM; if(r.ini)o.ini.push(toMin(r.ini)); if(r.fin)o.fin.push(toMin(r.fin)); if(r.hf!=null)o.hf.push(r.hf);});
  return Object.values(m).map(o=>{const tot=o.mL+o.mO+o.mM||1; return {...o,kmd:o.km/o.dias,hd:o.h/o.dias,pd:o.p/o.dias,pl:100*(o.mL+o.mO)/tot,pm:100*o.mM/tot,inim:avg(o.ini),finm:avg(o.fin),hfm:avg(o.hf),ex:D.excluidos.filter(x=>x.e===o.e&&x.f).length}}); }
function render(){
  const rows=D.diario.filter(keep), par=D.paradas.filter(keep); const pe=perEmp(rows);
  $('#fInfo').textContent=`${rows.length} dias-persona · ${pe.length} ejecutivos · ${par.length} paradas`;
  const tot=rows.reduce((a,r)=>a+r.mL+r.mO+r.mM,0)||1, km=rows.reduce((a,r)=>a+(r.km||0),0);
  const k=[['Ejecutivos',pe.length,''],['Dias con datos',rows.length,''],['Km recorridos',fmt(km,0),`${fmt(km/(rows.length||1))} km por dia`],['Jornada promedio',fmt(avg(rows.map(r=>r.h)))+' h','por dia con datos'],
    ['Inicio promedio',toHM(avg(rows.map(r=>toMin(r.ini)).filter(x=>x!=null))),'primer registro GPS'],['Fin promedio',toHM(avg(rows.map(r=>toMin(r.fin)).filter(x=>x!=null))),'ultimo registro GPS'],
    ['Paradas por dia',fmt(rows.reduce((a,r)=>a+(r.p||0),0)/(rows.length||1)),'de 10 min o mas'],['% tiempo detenido',fmt(100*rows.reduce((a,r)=>a+r.mL+r.mO,0)/tot,0)+'%','en aliados u otros lugares'],
    ['Horas fuera de base',fmt(avg(rows.map(r=>r.hf).filter(x=>x!=null)))+' h','salida a regreso al domicilio/base'],['Dias excluidos',D.excluidos.filter(x=>x.f&&(st.seg==='Todos'||D.segmentos[x.e]===st.seg)&&(!st.emp||x.e===st.emp)).length,'viajes / sin GPS']];
  $('#kpis').innerHTML=k.map(x=>`<div class="kpi"><div class="v">${x[1]}</div><div class="l">${x[0]}</div><div class="s">${x[2]}</div></div>`).join('');
  // segmentos
  $('#segcards').innerHTML=['Mercado Comercial','Residencial'].map(sg=>{const rs=D.diario.filter(r=>D.segmentos[r.e]===sg&&(!st.mes||r.f.startsWith(st.mes))&&(!st.ini||r.f>=st.ini)&&(!st.fin||r.f<=st.fin)); const t=rs.reduce((a,r)=>a+r.mL+r.mO+r.mM,0)||1; const n=new Set(rs.map(r=>r.e)).size;
    const row=(l,v)=>`<div class="row"><span>${l}</span><b>${v}</b></div>`; return `<div class="segcard ${sg==='Residencial'?'res':'com'}"><div style="font-weight:800;font-size:15px;margin-bottom:6px">${sg} <span class="pill">${n} ejecutivos · ${rs.length} dias</span></div>
    ${row('Km por dia',fmt(rs.reduce((a,r)=>a+(r.km||0),0)/(rs.length||1)))}${row('Jornada promedio',fmt(avg(rs.map(r=>r.h)))+' h')}${row('Inicio / fin promedio',toHM(avg(rs.map(r=>toMin(r.ini)).filter(x=>x!=null)))+' / '+toHM(avg(rs.map(r=>toMin(r.fin)).filter(x=>x!=null))))}
    ${row('Paradas por dia',fmt(rs.reduce((a,r)=>a+(r.p||0),0)/(rs.length||1)))}${row('% tiempo detenido / en movimiento',fmt(100*rs.reduce((a,r)=>a+r.mL+r.mO,0)/t,0)+'% / '+fmt(100*rs.reduce((a,r)=>a+r.mM,0)/t,0)+'%')}${row('Horas fuera de base',fmt(avg(rs.map(r=>r.hf).filter(x=>x!=null)))+' h')}
    ${row('Vel. maxima promedio',fmt(avg(rs.map(r=>r.vmax).filter(x=>x!=null)),0)+' km/h')}</div>`}).join('');
  // km por mes y dias por mes (por segmento)
  const byM={}; rows.forEach(r=>{const m=r.f.slice(0,7); const o=byM[m]||(byM[m]={}); const s=D.segmentos[r.e]; const q=o[s]||(o[s]={km:0,n:0,emps:new Set()}); q.km+=r.km||0;q.n++;q.emps.add(r.e)});
  const ms=Object.keys(byM).sort(); const segs=['Mercado Comercial','Residencial'].filter(s=>st.seg==='Todos'||st.seg===s);
  mk('cMes',{type:'bar',data:{labels:ms,datasets:segs.map(s=>({label:s+' · km/dia',data:ms.map(m=>byM[m][s]?+(byM[m][s].km/byM[m][s].n).toFixed(1):null),backgroundColor:SEGC[s]+'cc',borderRadius:6}))},options:{responsive:true,maintainAspectRatio:false,plugins:{legend:{position:'bottom'}},scales:{y:{beginAtZero:true,title:{display:true,text:'km por dia'}}}}});
  mk('cDias',{type:'line',data:{labels:ms,datasets:segs.map(s=>({label:s+' · dias-persona',data:ms.map(m=>byM[m][s]?byM[m][s].n:0),borderColor:SEGC[s],backgroundColor:SEGC[s]+'33',fill:true,tension:.3}))},options:{responsive:true,maintainAspectRatio:false,plugins:{legend:{position:'bottom'}},scales:{y:{beginAtZero:true}}}});
  // jornada floating bars
  const pj=[...pe].filter(o=>o.inim!=null).sort((a,b)=>a.inim-b.inim);
  mk('cJornada',{type:'bar',data:{labels:pj.map(o=>o.e),datasets:[{label:'Inicio → fin promedio',data:pj.map(o=>[o.inim/60,o.finm/60]),backgroundColor:pj.map(o=>SEGC[o.seg]+'cc'),borderRadius:6,borderSkipped:false}]},options:{indexAxis:'y',responsive:true,maintainAspectRatio:false,plugins:{legend:{display:false},tooltip:{callbacks:{label:c=>`${toHM(c.raw[0]*60)} → ${toHM(c.raw[1]*60)} (${fmt(c.raw[1]-c.raw[0])} h)`}}},scales:{x:{min:5,max:21,ticks:{callback:v=>String(v).padStart(2,'0')+':00'}}}}});
  // tiempo stacked
  const pt=[...pe].sort((a,b)=>b.pl-a.pl);
  mk('cTiempo',{type:'bar',data:{labels:pt.map(o=>o.e),datasets:[{label:'% detenido (aliados u otros lugares)',data:pt.map(o=>+o.pl.toFixed(1)),backgroundColor:'#f59e0bcc',borderRadius:4},{label:'% en movimiento',data:pt.map(o=>+o.pm.toFixed(1)),backgroundColor:'#38bdf8cc',borderRadius:4}]},options:{indexAxis:'y',responsive:true,maintainAspectRatio:false,plugins:{legend:{position:'bottom'}},scales:{x:{stacked:true,max:100},y:{stacked:true}}}});
  // ranking
  const key={e:o=>o.e,seg:o=>o.seg,dias:o=>o.dias,kmd:o=>o.kmd,km:o=>o.km,ini:o=>o.inim??1e9,fin:o=>o.finm??1e9,h:o=>o.hd,pd:o=>o.pd,pl:o=>o.pl,pm:o=>o.pm,hf:o=>o.hfm??-1,ex:o=>o.ex}[st.sortK];
  const pr=[...pe].sort((a,b)=>{const x=key(a),y=key(b); return (x>y?1:x<y?-1:0)*st.sortD}); const maxkm=Math.max(...pe.map(o=>o.kmd),1);
  $('#rank tbody').innerHTML=pr.map(o=>`<tr data-e="${o.e}" ${st.det===o.e?'style="background:rgba(99,102,241,.12)"':''}><td><b>${o.e}</b></td><td><span class="badge ${o.seg==='Residencial'?'b-res':'b-com'}">${o.seg==='Residencial'?'Residencial':'Comercial'}</span></td><td class="r">${o.dias}</td><td class="r">${fmt(o.kmd)}<span class="bar" style="width:${60*o.kmd/maxkm}px"></span></td><td class="r">${fmt(o.km,0)}</td><td class="r">${toHM(o.inim)}</td><td class="r">${toHM(o.finm)}</td><td class="r">${fmt(o.hd)}</td><td class="r">${fmt(o.pd)}</td><td class="r">${fmt(o.pl,0)}%</td><td class="r">${fmt(o.pm,0)}%</td><td class="r">${o.hfm==null?'—':fmt(o.hfm)}</td><td class="r">${o.ex}</td></tr>`).join('');
  document.querySelectorAll('#rank tbody tr').forEach(tr=>tr.onclick=()=>{st.det=tr.dataset.e;st.day='';renderDetail();document.querySelectorAll('#rank tbody tr').forEach(x=>x.style.background='');tr.style.background='rgba(99,102,241,.12)'});
  // heatmap dia x hora
  const days=['Lun','Mar','Mie','Jue','Vie']; const hours=[...Array(14)].map((_,i)=>i+6); const H={}; let hm=0;
  par.forEach(p=>{const d=new Date(p.f+'T00:00:00').getDay()-1; const h=parseInt(p.ll.slice(0,2)); if(d<0||d>4||h<6||h>19)return; H[d+'-'+h]=(H[d+'-'+h]||0)+1; hm=Math.max(hm,H[d+'-'+h])});
  $('#heat').innerHTML='<div></div>'+hours.map(h=>`<div style="color:var(--mu)">${h}h</div>`).join('')+days.map((d,i)=>`<div style="text-align:left;color:var(--mu)">${d}</div>`+hours.map(h=>{const v=H[i+'-'+h]||0; const a=hm?v/hm:0; return `<div title="${d} ${h}:00 · ${v} paradas" style="background:rgba(99,102,241,${(0.08+a*0.9).toFixed(2)})">${v||''}</div>`}).join('')).join('');
  // duracion de paradas
  const bins=[['10-20 min',10,20],['20-40',20,40],['40-60',40,60],['1-2 h',60,120],['2-4 h',120,240],['> 4 h',240,1e9]];
  mk('cDur',{type:'bar',data:{labels:bins.map(b=>b[0]),datasets:[{label:'Paradas',data:bins.map(b=>par.filter(p=>p.min>=b[1]&&p.min<b[2]).length),backgroundColor:'#22d3eecc',borderRadius:6}]},options:{responsive:true,maintainAspectRatio:false,plugins:{legend:{display:false}}}});
  // lugares + mapa
  const lug=D.lugares.filter(l=>(st.seg==='Todos'||D.segmentos[l.e]===st.seg)&&(!st.emp||l.e===st.emp)).sort((a,b)=>b.mt-a.mt);
  $('#tLug tbody').innerHTML=lug.slice(0,80).map(l=>`<tr><td><b>${l.e}</b><div class="note">${(l.dir||'').slice(0,48)}</div></td><td>${l.t.startsWith('Posible')?'<span style="color:#fca5a5">Base/domicilio</span>':l.t.startsWith('Aliado')?'<span style="color:#6ee7b7">'+l.t+'</span>':'Recurrente'}</td><td class="r">${l.v}</td><td class="r">${fmt(l.mt,0)}</td><td>${l.ll||''}</td></tr>`).join('');
  if(!window._map){window._map=L.map('map').setView([4.7,-74.1],7);L.tileLayer('https://{s}.basemaps.cartocdn.com/rastertiles/voyager/{z}/{x}/{y}{r}.png',{maxZoom:19}).addTo(window._map);window._lay=L.layerGroup().addTo(window._map)}
  window._lay.clearLayers(); const pts=[];
  lug.forEach(l=>{if(l.lat==null)return; const col=l.t.startsWith('Posible')?'#ef4444':SEGC[D.segmentos[l.e]]; const r=6+Math.min(22,Math.sqrt(l.mt||1)/3);
    L.circleMarker([l.lat,l.lng],{radius:r,color:'#fff',weight:1,fillColor:col,fillOpacity:.8}).addTo(window._lay).bindPopup(`<b>${l.e}</b><br>${l.t}<br>${l.v} veces · ${fmt(l.mt,0)} min · llega ${l.ll||''}<br><small>${l.dir||''}</small>`); pts.push([l.lat,l.lng])});
  if(pts.length)window._map.fitBounds(pts,{padding:[20,20]});
  $('#tEx tbody').innerHTML=D.excluidos.filter(x=>(st.seg==='Todos'||D.segmentos[x.e]===st.seg)&&(!st.emp||x.e===st.emp)).map(x=>`<tr><td>${x.e}</td><td>${x.f||''}</td><td>${x.m}</td></tr>`).join('');
  if(st.emp&&st.det!==st.emp){st.det=st.emp;st.day=''} renderDetail();
}
function renderDetail(){
  const e=st.det; if(!e){$('#detName').textContent='(selecciona uno en el ranking o en el filtro)';$('#cal').innerHTML='';return}
  $('#detName').textContent=e+' · '+D.segmentos[e];
  const rows=D.diario.filter(r=>r.e===e&&(!st.mes||r.f.startsWith(st.mes))&&(!st.ini||r.f>=st.ini)&&(!st.fin||r.f<=st.fin)).sort((a,b)=>a.f<b.f?-1:1); const mx=Math.max(...rows.map(r=>r.km||0),1);
  $('#cal').innerHTML=rows.map(r=>`<div data-f="${r.f}" title="${r.f} · ${fmt(r.km)} km · ${r.ini}-${r.fin} · ${r.p} paradas" style="cursor:pointer;background:rgba(99,102,241,${(0.1+0.85*(r.km||0)/mx).toFixed(2)});${st.day===r.f?'outline:2px solid #fbbf24':''}"><div style="font-size:9px;color:#cbd5e1">${r.f.slice(5)}</div><b>${fmt(r.km,0)}</b><div style="font-size:9px">${r.ini}</div></div>`).join('');
  document.querySelectorAll('#cal div[data-f]').forEach(d=>d.onclick=()=>{st.day=d.dataset.f;renderDetail()});
  if(!st.day){$('#detDay').textContent='';$('#tl').innerHTML='<div class="note">Elige un dia en el calendario.</div>';return}
  const r=rows.find(x=>x.f===st.day); const ps=D.paradas.filter(p=>p.e===e&&p.f===st.day).sort((a,b)=>a.n-b.n);
  $('#detDay').textContent=st.day+(r?` · ${r.ini}–${r.fin} · ${fmt(r.km)} km`:'');
  let html=`<div class="it"><div class="t">${r?r.ini:''}</div><div class="box">Primer registro GPS del dia${r&&r.sb?` · sale de base ${r.sb}`:''}</div></div>`;
  ps.forEach(p=>{html+=`<div class="mv">⟶ ${fmt(p.tr,0)} min en desplazamiento${p.kma?` · ${fmt(p.kma)} km`:''}</div><div class="it"><div class="t">${p.ll}</div><div class="box"><b>${p.t==='Aliado'?'Aliado: '+p.al:(p.lf||'Parada')}</b> · ${fmt(p.min,0)} min (hasta ${p.sa})<div class="note">${p.dir||''} <a href="https://www.google.com/maps?q=${p.lat},${p.lng}" target="_blank" style="color:#818cf8">mapa</a></div></div></div>`});
  html+=`<div class="it"><div class="t">${r?r.fin:''}</div><div class="box">Ultimo registro GPS${r&&r.lb?` · llega a base ${r.lb}`:''}</div></div>`; $('#tl').innerHTML=html;
}
document.querySelectorAll('#rank th').forEach(th=>th.onclick=()=>{const k=th.dataset.k; if(st.sortK===k)st.sortD*=-1; else {st.sortK=k;st.sortD=(k==='e'||k==='seg'||k==='ini')?1:-1} render()});
render();
</script></body></html>'''

from datetime import datetime
html = HTML.replace('__DATA__', json.dumps(data, ensure_ascii=False, separators=(',', ':'))).replace('__FECHA__', datetime.now().strftime('%Y-%m-%d %H:%M'))
out_html = os.path.join(os.path.dirname(XLSX), 'VantiGo_Super_Dashboard_Tiempos.html')
io.open(out_html, 'w', encoding='utf8').write(html)
print('HTML ->', out_html, f'({len(html)//1024} KB)')

# ---------- Hoja "Dashboard" con graficos nativos dentro del Excel ----------
wb2 = load_workbook(XLSX)
if 'Dashboard' in wb2.sheetnames:
    del wb2['Dashboard']
ws = wb2.create_sheet('Dashboard', 0)
hf = PatternFill('solid', fgColor='4F46E5'); hfont = Font(bold=True, color='FFFFFF')
ws['A1'] = 'SUPER DASHBOARD · Uso del tiempo de los ejecutivos (06:00-20:00, L-V)'; ws['A1'].font = Font(bold=True, size=14, color='4F46E5')
ws['A2'] = f"Periodo {data['periodo'][0]} a {data['periodo'][1]} · Mercado Comercial = registro antiguo · Residencial = bloque del 20-27 de agosto · Version interactiva: VantiGo_Super_Dashboard_Tiempos.html"
# tabla por ejecutivo
agg = defaultdict(lambda: {'dias': 0, 'km': 0, 'h': 0, 'p': 0, 'det': 0, 'mov': 0, 'ini': [], 'fin': [], 'hf': []})
for r in data['diario']:
    a = agg[r['e']]; a['dias'] += 1; a['km'] += r['km'] or 0; a['h'] += r['h'] or 0; a['p'] += r['p'] or 0; a['det'] += r['mL'] + r['mO']; a['mov'] += r['mM']
    if r['ini']: a['ini'].append(int(r['ini'][:2]) * 60 + int(r['ini'][3:5]))
    if r['fin']: a['fin'].append(int(r['fin'][:2]) * 60 + int(r['fin'][3:5]))
    if r['hf'] is not None: a['hf'].append(r['hf'])
hm = lambda l: (f"{int(sum(l)/len(l))//60:02d}:{int(sum(l)/len(l))%60:02d}" if l else '')
hdr = ['Ejecutivo', 'Segmento', 'Dias', 'Km total', 'Km/dia', 'Inicio prom.', 'Fin prom.', 'Jornada (h)', 'Paradas/dia', '% detenido', '% movimiento', 'Horas fuera de base']
ws.append([]); ws.append(hdr)
for c in ws[4]: c.fill = hf; c.font = hfont
rows_e = sorted(agg.items(), key=lambda kv: -(kv[1]['km'] / (kv[1]['dias'] or 1)))
for e, a in rows_e:
    d = a['dias'] or 1; t = (a['det'] + a['mov']) or 1
    ws.append([e, SEG(e), a['dias'], round(a['km'], 1), round(a['km'] / d, 1), hm(a['ini']), hm(a['fin']), round(a['h'] / d, 2), round(a['p'] / d, 1), round(100 * a['det'] / t, 1), round(100 * a['mov'] / t, 1), round(sum(a['hf']) / len(a['hf']), 2) if a['hf'] else ''])
last = 4 + len(rows_e)
# tabla por segmento
r0 = last + 3
ws.cell(row=r0, column=1, value='Comparativo por segmento').font = Font(bold=True, color='4F46E5')
ws.append([]); hdr2 = ['Segmento', 'Ejecutivos', 'Dias', 'Km/dia', 'Jornada (h)', 'Inicio prom.', 'Fin prom.', 'Paradas/dia', '% detenido', '% movimiento']
ws.append(hdr2)
for c in ws[ws.max_row]: c.fill = hf; c.font = hfont
for sg in ('Mercado Comercial', 'Residencial'):
    rs = [r for r in data['diario'] if SEG(r['e']) == sg]
    if not rs: continue
    t = sum(r['mL'] + r['mO'] + r['mM'] for r in rs) or 1
    ws.append([sg, len({r['e'] for r in rs}), len(rs), round(sum(r['km'] or 0 for r in rs) / len(rs), 1), round(sum(r['h'] or 0 for r in rs) / len(rs), 2),
               hm([int(r['ini'][:2]) * 60 + int(r['ini'][3:5]) for r in rs if r['ini']]), hm([int(r['fin'][:2]) * 60 + int(r['fin'][3:5]) for r in rs if r['fin']]),
               round(sum(r['p'] or 0 for r in rs) / len(rs), 1), round(100 * sum(r['mL'] + r['mO'] for r in rs) / t, 1), round(100 * sum(r['mM'] for r in rs) / t, 1)])
# tabla por mes y segmento
r1 = ws.max_row + 3
ws.cell(row=r1, column=1, value='Km por dia y dias con actividad por mes').font = Font(bold=True, color='4F46E5')
ws.append([]); ws.append(['Mes', 'Km/dia Comercial', 'Km/dia Residencial', 'Dias Comercial', 'Dias Residencial'])
for c in ws[ws.max_row]: c.fill = hf; c.font = hfont
first_m = ws.max_row + 1
mon = defaultdict(lambda: {'Mercado Comercial': [0, 0], 'Residencial': [0, 0]})
for r in data['diario']:
    m = mon[r['f'][:7]][SEG(r['e'])]; m[0] += r['km'] or 0; m[1] += 1
for mth in sorted(mon):
    c, rr = mon[mth]['Mercado Comercial'], mon[mth]['Residencial']
    ws.append([mth, round(c[0] / c[1], 1) if c[1] else None, round(rr[0] / rr[1], 1) if rr[1] else None, c[1], rr[1]])
last_m = ws.max_row
# graficos
ch = BarChart(); ch.type = 'bar'; ch.title = 'Km por dia por ejecutivo'; ch.y_axis.title = 'km/dia'; ch.height = 9; ch.width = 18
ch.add_data(Reference(ws, min_col=5, min_row=4, max_row=last), titles_from_data=True); ch.set_categories(Reference(ws, min_col=1, min_row=5, max_row=last)); ws.add_chart(ch, 'N4')
ch2 = BarChart(); ch2.type = 'bar'; ch2.grouping = 'percentStacked'; ch2.overlap = 100; ch2.title = '% detenido vs en movimiento'; ch2.height = 9; ch2.width = 18
ch2.add_data(Reference(ws, min_col=10, max_col=11, min_row=4, max_row=last), titles_from_data=True); ch2.set_categories(Reference(ws, min_col=1, min_row=5, max_row=last)); ws.add_chart(ch2, 'N24')
ch3 = LineChart(); ch3.title = 'Km por dia por mes y segmento'; ch3.height = 8; ch3.width = 18
ch3.add_data(Reference(ws, min_col=2, max_col=3, min_row=first_m - 1, max_row=last_m), titles_from_data=True); ch3.set_categories(Reference(ws, min_col=1, min_row=first_m, max_row=last_m)); ws.add_chart(ch3, 'N44')
ch4 = BarChart(); ch4.title = 'Dias con actividad por mes'; ch4.height = 8; ch4.width = 18
ch4.add_data(Reference(ws, min_col=4, max_col=5, min_row=first_m - 1, max_row=last_m), titles_from_data=True); ch4.set_categories(Reference(ws, min_col=1, min_row=first_m, max_row=last_m)); ws.add_chart(ch4, 'N62')
for col, w in zip('ABCDEFGHIJKL', (16, 18, 7, 10, 9, 12, 11, 11, 11, 11, 12, 18)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A5'
out_xlsx = XLSX.replace('.xlsx', '_con_Dashboard.xlsx')
wb2.save(out_xlsx)
print('XLSX ->', out_xlsx)
