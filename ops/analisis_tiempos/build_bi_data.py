# -*- coding: utf-8 -*-
"""Paso 1 del Excel BI: prepara las TABLAS de datos (ListObjects) que alimentan las tablas
dinamicas, segmentadores y graficos que crea build_bi_excel.ps1 con Excel.
Uso: python build_bi_data.py <Analisis..._Geoanalitica.xlsx> <salida_base.xlsx>"""
import math, sys, statistics
from collections import defaultdict
from datetime import date, datetime
from openpyxl import load_workbook, Workbook
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SRC, OUT = sys.argv[1], sys.argv[2]
RESIDENCIAL = {'amgomez', 'jjbecerra', 'kvlamprea', 'vqherrera', 'gmalvarado', 'elmateus', 'lemurcia', 'dramaya', 'malejandra', 'kmlara', 'ajsuarez'}
SEG = lambda e: 'Residencial' if e in RESIDENCIAL else 'Mercado Comercial'
hav = lambda a, b, c, d: 2 * 6371000 * math.asin(math.sqrt(math.sin(math.radians(c - a) / 2) ** 2 + math.cos(math.radians(a)) * math.cos(math.radians(c)) * math.sin(math.radians(d - b) / 2) ** 2))
num = lambda v: (float(v) if v not in (None, '') else None)
DIAS = ['Lunes', 'Martes', 'Miercoles', 'Jueves', 'Viernes', 'Sabado', 'Domingo']
MESES = ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic']

wb = load_workbook(SRC, read_only=True)
def sheet(n):
    it = wb[n].iter_rows(values_only=True); h = next(it); return [dict(zip(h, r)) for r in it if any(v not in (None, '') for v in r)]
diario, paradas, lugares, kpis, concl, excl = sheet('Resumen diario'), sheet('Paradas'), sheet('Lugares frecuentes'), sheet('KPIs geo-operativos'), sheet('Conclusiones'), sheet('Dias excluidos')
wb.close()
bases = {l['Empleado']: (num(l['Latitud']), num(l['Longitud'])) for l in lugares if str(l.get('Tipo') or l.get('Tipo de lugar') or '').startswith('Posible')}

def tipo_parada(p):
    b = bases.get(p['Empleado']); d = num(p['Duracion (min)']) or 0
    if b and hav(num(p['Latitud']), num(p['Longitud']), b[0], b[1]) <= 150: return 'En base/domicilio'
    if d >= 120: return 'Tiempo muerto (>=2h)'
    if d >= 20: return 'Productiva (>=20 min)'
    return 'Corta (<20 min)'
for p in paradas: p['_tipo'] = tipo_parada(p)
pday = defaultdict(lambda: defaultdict(int)); pday_min = defaultdict(lambda: defaultdict(float))
for p in paradas:
    k = (p['Empleado'], p['Fecha']); pday[k][p['_tipo']] += 1; pday_min[k][p['_tipo']] += num(p['Duracion (min)']) or 0

out = Workbook(); out.remove(out.active)
def mk_table(ws, name, headers, rows, style='TableStyleMedium2', widths=None, date_cols=()):
    ws.append(headers)
    for r in rows: ws.append(r)
    for i, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, min(max(9, len(str(h)) + 2), 40))
        if h in date_cols:
            for row in range(2, len(rows) + 2): ws.cell(row=row, column=i).number_format = 'yyyy-mm-dd'
    ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    t = Table(displayName=name, ref=ref); t.tableStyleInfo = TableStyleInfo(name=style, showRowStripes=True); ws.add_table(t); ws.freeze_panes = 'A2'

# --- Diario ---
H = ['Empleado', 'Segmento', 'Fecha', 'Mes', 'Semana', 'Dia semana', 'Inicio', 'Fin', 'Jornada h', 'Km', 'Paradas', 'Paradas productivas', 'Tiempos muertos', 'Min en base', 'Min productivos', 'Min tiempo muerto', 'Min detenido', 'Min movimiento', 'Vmax', 'Exceso >80', 'Salida base', 'Llegada base', 'Horas fuera base']
rows = []
for r in diario:
    e, f = r['Empleado'], r['Fecha']; d = date.fromisoformat(f); k = (e, f); vmax = num(r['Vel. max (km/h)']) or 0
    rows.append([e, SEG(e), d, f"{d.year}-{d.month:02d} {MESES[d.month - 1]}", d.isocalendar()[1], DIAS[d.weekday()], r['Inicio recorrido'], r['Fin recorrido'], num(r['Jornada (h)']), num(r['Km recorridos']), num(r['Paradas']) or 0,
                 pday[k]['Productiva (>=20 min)'], pday[k]['Tiempo muerto (>=2h)'], round(pday_min[k]['En base/domicilio']), round(pday_min[k]['Productiva (>=20 min)']), round(pday_min[k]['Tiempo muerto (>=2h)']),
                 (num(r['Min en aliados']) or 0) + (num(r['Min en otros lugares']) or 0), num(r['Min en movimiento']) or 0, vmax if vmax <= 200 else None, 1 if 80 < vmax <= 200 else 0,
                 r.get('Salida de base/domicilio') or '', r.get('Llegada a base/domicilio') or '', num(r.get('Horas fuera de base'))])
rows.sort(key=lambda x: (x[2], x[0]))
mk_table(out.create_sheet('Datos_Diario'), 'tblDiario', H, rows, date_cols=('Fecha',), widths={'Empleado': 14, 'Segmento': 17, 'Fecha': 11})
# --- Paradas ---
H2 = ['Empleado', 'Segmento', 'Fecha', 'Mes', 'Dia semana', 'Hora llegada', 'Llegada', 'Salida', 'Duracion min', 'Duracion h', 'Tipo parada', 'Lugar frecuente', 'Direccion', 'Latitud', 'Longitud', 'Mapa']
rows2 = []
for p in paradas:
    d = date.fromisoformat(p['Fecha']); dm = num(p['Duracion (min)']) or 0
    rows2.append([p['Empleado'], SEG(p['Empleado']), d, f"{d.year}-{d.month:02d} {MESES[d.month - 1]}", DIAS[d.weekday()], int(p['Llegada'][:2]) if p['Llegada'] else None, p['Llegada'], p['Salida'], dm, round(dm / 60, 2), p['_tipo'], p.get('Lugar frecuente') or '', (p.get('Direccion') or '')[:80], num(p['Latitud']), num(p['Longitud']), f"https://www.google.com/maps?q={p['Latitud']},{p['Longitud']}"])
mk_table(out.create_sheet('Datos_Paradas'), 'tblParadas', H2, rows2, 'TableStyleMedium6', date_cols=('Fecha',), widths={'Direccion': 44, 'Mapa': 40, 'Tipo parada': 22, 'Lugar frecuente': 30})
# --- KPIs por ejecutivo ---
H3 = list(kpis[0].keys())
rows3 = [[k[h] for h in H3] for k in kpis]
mk_table(out.create_sheet('Datos_KPI'), 'tblKPI', H3, rows3, 'TableStyleMedium9')
# --- Lugares ---
H4 = ['Empleado', 'Segmento', 'Lugar', 'Tipo', 'Veces', 'Dias distintos', 'Llegada promedio', 'Duracion promedio min', 'Minutos totales', 'Direccion', 'Latitud', 'Longitud', 'Mapa']
rows4 = [[l['Empleado'], SEG(l['Empleado']), l['Lugar'], str(l.get('Tipo') or l.get('Tipo de lugar') or ''), num(l['Veces']), num(l['Dias distintos']), l['Llegada promedio'], num(l['Duracion promedio (min)']), num(l['Minutos totales']), (l.get('Direccion') or '')[:80], num(l['Latitud']), num(l['Longitud']), f"https://www.google.com/maps?q={l['Latitud']},{l['Longitud']}"] for l in lugares]
mk_table(out.create_sheet('Datos_Lugares'), 'tblLugares', H4, rows4, 'TableStyleMedium4', widths={'Direccion': 44, 'Mapa': 40})
# --- Conclusiones y excluidos ---
ws = out.create_sheet('Conclusiones')
ws.append(['#', 'Conclusion (con cifras)', 'Indicador', 'Donde verlo'])
for c in ws[1]: c.fill = PatternFill('solid', fgColor='4F46E5'); c.font = Font(bold=True, color='FFFFFF')
for c in concl: ws.append([c['#'], c['Conclusion (con cifras)'], c.get('Indicador que la sustenta', ''), c.get('Hoja donde verlo', '')])
ws.column_dimensions['B'].width = 120; ws.column_dimensions['C'].width = 36; ws.column_dimensions['D'].width = 24
for row in ws.iter_rows(min_row=2): row[1].alignment = Alignment(wrap_text=True, vertical='top')
mk_table(out.create_sheet('Datos_Excluidos'), 'tblExcluidos', ['Empleado', 'Segmento', 'Fecha', 'Motivo'], [[x['Empleado'], SEG(x['Empleado']), x['Fecha'], x['Motivo']] for x in excl if x['Fecha']], 'TableStyleLight9', widths={'Motivo': 70})
out.save(OUT)
print('OK ->', OUT, '| diario', len(rows), '| paradas', len(rows2), '| kpi', len(rows3), '| lugares', len(rows4))
